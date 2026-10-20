"""Generación de reportes PDF/Word/Excel con límite de tamaño y expiración."""
import logging
import os
import uuid
from datetime import datetime, timedelta
from pathlib import Path

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

from docx import Document as DocxDocument
from docx.shared import Pt

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from app.config import settings
from app.services import metrics

log = logging.getLogger(__name__)

EXPORT_DIR = Path(settings.export_dir)
EXPORT_DIR.mkdir(parents=True, exist_ok=True)


def _filename(prefix: str, ext: str) -> Path:
    stamp = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
    return EXPORT_DIR / f"{prefix}_{stamp}_{uuid.uuid4().hex[:6]}.{ext}"


def _enforce_size_limit(path: Path) -> None:
    size_mb = path.stat().st_size / (1024 * 1024)
    if size_mb > settings.export_max_size_mb:
        path.unlink(missing_ok=True)
        raise RuntimeError(f"Reporte excede {settings.export_max_size_mb} MB")


def _rows_from_summary(s: dict):
    yield ["Métrica", "Valor"]
    yield ["Consultas totales", s["total_queries"]]
    yield ["Resueltas sin escalar", s["resolved_without_escalation"]]
    yield ["Incidentes reportados", s["incidents_reported"]]
    yield ["Tasa de escalamiento (%)", s["escalation_rate"]]
    yield ["Latencia promedio (ms)", s["avg_latency_ms"]]
    yield ["Latencia p95 (ms)", s["p95_latency_ms"]]
    yield ["Cobertura de citas (%)", s["citation_coverage"]]
    yield ["Horas ahorradas", s["hours_saved"]]


def _topics_table(s: dict):
    yield ["Tema", "Consultas"]
    for t in s["top_topics"]:
        yield [t["topic"], t["n"]]


def _tickets_status(s: dict):
    yield ["Estado", "Cantidad"]
    for k, v in s["tickets_by_status"].items():
        yield [k, v]


def _tickets_severity(s: dict):
    yield ["Severidad", "Cantidad"]
    for k, v in s["tickets_by_severity"].items():
        yield [k, v]


def _queries_by_day(s: dict):
    yield ["Día", "Consultas"]
    for d in s["queries_by_day"]:
        yield [d["day"], d["count"]]


def _pdf_table(data):
    t = Table(data, hAlign="LEFT")
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f4e79")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f2f2f2")]),
    ]))
    return t


def export_pdf(report_type: str, s: dict) -> Path:
    path = _filename(f"report_{report_type}", "pdf")
    styles = getSampleStyleSheet()
    story = [
        Paragraph(f"Reporte {report_type}", styles["Title"]),
        Paragraph(f"Generado: {datetime.utcnow():%Y-%m-%d %H:%M UTC}", styles["Normal"]),
        Spacer(1, 12),
    ]
    for title, gen in [
        ("Resumen", _rows_from_summary),
        ("Temas frecuentes", _topics_table),
        ("Tickets por estado", _tickets_status),
        ("Tickets por severidad", _tickets_severity),
        ("Consultas por día", _queries_by_day),
    ]:
        story.append(Paragraph(title, styles["Heading2"]))
        story.append(_pdf_table(list(gen(s))))
        story.append(Spacer(1, 12))
    SimpleDocTemplate(str(path), pagesize=A4).build(story)
    _enforce_size_limit(path)
    return path


def export_docx(report_type: str, s: dict) -> Path:
    path = _filename(f"report_{report_type}", "docx")
    doc = DocxDocument()
    doc.add_heading(f"Reporte {report_type}", level=0)
    doc.add_paragraph(f"Generado: {datetime.utcnow():%Y-%m-%d %H:%M UTC}")

    def add_table(title, rows):
        doc.add_heading(title, level=1)
        rows = list(rows)
        table = doc.add_table(rows=0, cols=len(rows[0]))
        table.style = "Light Grid Accent 1"
        for i, row in enumerate(rows):
            cells = table.add_row().cells
            for j, val in enumerate(row):
                cells[j].text = str(val)
                if i == 0:
                    for p in cells[j].paragraphs:
                        for r in p.runs:
                            r.bold = True
                            r.font.size = Pt(11)

    add_table("Resumen", _rows_from_summary(s))
    add_table("Temas frecuentes", _topics_table(s))
    add_table("Tickets por estado", _tickets_status(s))
    add_table("Tickets por severidad", _tickets_severity(s))
    add_table("Consultas por día", _queries_by_day(s))
    doc.save(str(path))
    _enforce_size_limit(path)
    return path


def export_xlsx(report_type: str, s: dict) -> Path:
    path = _filename(f"report_{report_type}", "xlsx")
    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E79")

    def add_sheet(name, rows):
        ws = wb.create_sheet(title=name[:31])
        rows = list(rows)
        for row in rows:
            ws.append(row)
        for col in range(1, len(rows[0]) + 1):
            c = ws.cell(row=1, column=col)
            c.font = header_font
            c.fill = header_fill
            ws.column_dimensions[chr(64 + col)].width = 32

    wb.remove(wb.active)
    add_sheet("Resumen", _rows_from_summary(s))
    add_sheet("Temas", _topics_table(s))
    add_sheet("EstadoTickets", _tickets_status(s))
    add_sheet("SeveridadTickets", _tickets_severity(s))
    add_sheet("ConsultasPorDia", _queries_by_day(s))
    wb.save(str(path))
    _enforce_size_limit(path)
    return path


def cleanup_expired() -> int:
    """Elimina archivos expirados del directorio de exportación."""
    cutoff = datetime.utcnow() - timedelta(days=settings.export_max_age_days)
    removed = 0
    for f in EXPORT_DIR.iterdir():
        try:
            if f.is_file() and datetime.utcfromtimestamp(f.stat().st_mtime) < cutoff:
                f.unlink()
                removed += 1
        except Exception as e:
            log.warning("cleanup_failed", extra={"file": str(f), "err": str(e)})
    return removed


EXPORTERS = {"pdf": export_pdf, "docx": export_docx, "xlsx": export_xlsx}